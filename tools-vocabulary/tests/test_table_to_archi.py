import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path
import xml.etree.ElementTree as ET

import openpyxl


MODULE_PATH = Path(__file__).resolve().parents[1] / "tableToArchi.py"
SPEC = importlib.util.spec_from_file_location("tableToArchi", MODULE_PATH)
MODULE = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = MODULE
SPEC.loader.exec_module(MODULE)


class TableToArchiViewTests(unittest.TestCase):
    def test_class_overview_contains_specialization_connections(self):
        workbook = openpyxl.Workbook()
        vocabulary = workbook.active
        vocabulary.title = MODULE.VOCABULARY
        vocabulary.append(["Název", "Test vocabulary"])

        classes = workbook.create_sheet(MODULE.CLASSES)
        classes.append([MODULE.NAME, MODULE.TYPE, MODULE.PARENT])
        classes.append(["Parent", "objekt práva", None])
        classes.append(["Child", "objekt práva", "Parent"])

        tropes = workbook.create_sheet(MODULE.TROPES)
        tropes.append([MODULE.NAME, MODULE.ENDPOINT])

        relationships = workbook.create_sheet(MODULE.RELATIONSHIPS)
        relationships.append([MODULE.NAME, MODULE.ENDPOINT, MODULE.ENDPOINT])
        relationships.append(["related to", "Parent", "Child"])

        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "source.xlsx"
            output = Path(temp) / "model.xml"
            workbook.save(source)
            MODULE.convert(source, output, with_view=True)
            root = ET.parse(output).getroot()

        ns = {"a": MODULE.NS}
        xsi_type = f"{{{MODULE.XSI}}}type"
        model_links = root.findall("a:relationships/a:relationship", ns)
        expected = {link.attrib["identifier"] for link in model_links
                    if link.attrib[xsi_type] in ("Association", "Specialization")}
        overview = root.find("a:views/a:diagrams/a:view[@identifier='id-class-overview']", ns)
        connections = overview.findall("a:connection", ns)
        self.assertEqual({connection.attrib["relationshipRef"] for connection in connections}, expected)
        node_ids = {node.attrib["identifier"] for node in overview.findall("a:node", ns)}
        self.assertTrue(all(connection.attrib["source"] in node_ids
                            and connection.attrib["target"] in node_ids
                            for connection in connections))
        nodes = {node.attrib["identifier"]: node for node in overview.findall("a:node", ns)}
        top_y = min(int(node.attrib["y"]) for node in nodes.values())
        self.assertTrue(all(int(point.attrib["y"]) >= top_y
                            for connection in connections
                            for point in connection.findall("a:bendpoint", ns)))
        specialization = next(link for link in model_links
                              if link.attrib[xsi_type] == "Specialization")
        connection = next(connection for connection in connections
                          if connection.attrib["relationshipRef"] == specialization.attrib["identifier"])
        points = connection.findall("a:bendpoint", ns)
        self.assertEqual(len(points), 4)
        self.assertEqual(int(points[0].attrib["y"]), int(nodes[connection.attrib["source"]].attrib["y"]))
        target = nodes[connection.attrib["target"]]
        self.assertEqual(int(points[-1].attrib["y"]),
                         int(target.attrib["y"]) + int(target.attrib["h"]))

    def test_same_row_association_stays_below_nodes(self):
        diagrams = ET.Element(f"{{{MODULE.NS}}}diagrams")
        link = ET.Element(f"{{{MODULE.NS}}}relationship", {
            "identifier": "association", "source": "left", "target": "right",
            f"{{{MODULE.XSI}}}type": "Association",
        })
        MODULE.add_hierarchical_view(diagrams, "overview", "Overview",
                                     ["left", "right"], [link],
                                     layout_relationships=[])
        node = diagrams.find(f"{{{MODULE.NS}}}view/{{{MODULE.NS}}}node")
        connection = diagrams.find(f"{{{MODULE.NS}}}view/{{{MODULE.NS}}}connection")
        points = connection.findall(f"{{{MODULE.NS}}}bendpoint")
        bottom = int(node.attrib["y"]) + int(node.attrib["h"])
        self.assertEqual(len(points), 4)
        self.assertTrue(all(int(point.attrib["y"]) >= bottom for point in points))


if __name__ == "__main__":
    unittest.main()
