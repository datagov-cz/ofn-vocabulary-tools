#!/usr/bin/env python3
"""Convert a tableToOFN Excel workbook to ArchiMate 3.1 exchange XML.

Usage: python tableToArchi.py input.xlsx output.xml [--with-view]
Requires Python 3.9+ and openpyxl; usable outside this repository.
Reads Slovník, Subjekty a objekty práva, Vlastnosti, and Vztahy worksheets.

Classes and properties become BusinessObjects, property ownership becomes
Composition, relationships become directed Associations, and parent concepts
become Specializations. All nonempty metadata columns are kept as properties.
References may be case-insensitive names or exact IRIs. External IRIs get
placeholder BusinessObjects (may also be exported as terms by archiToOFN.py).

Importer limitations: archiToOFN.py ignores the catalogue address and cannot
represent an untyped class as a TermClass. These cases produce warnings.
With --with-view, create a subject/object overview and a property view for each
owner with properties. Views use top-down layers with at most ten nodes per row. Formula cells use Excel's last saved values.
"""

import argparse
from dataclasses import dataclass
from pathlib import Path
import re
import sys
from urllib.parse import unquote
import warnings
from zipfile import BadZipFile
import xml.etree.ElementTree as ET

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException


NS = "http://www.opengroup.org/xsd/archimate/3.0/"
XSI = "http://www.w3.org/2001/XMLSchema-instance"
LANG = {"{http://www.w3.org/XML/1998/namespace}lang": "cs"}
VOCABULARY = "Slovník"
CLASSES = "Subjekty a objekty práva"
TROPES = "Vlastnosti"
RELATIONSHIPS = "Vztahy"
NAME = "Název"
TYPE = "Typ"
IRI = "Identifikátor"
PARENT = "Nadřazený pojem"
ENDPOINT = "Subjekt nebo objekt práva"
CATALOGUE = "Adresa lokálního katalogu dat, ve kterém bude slovník registrován"


def text(value):
    return "" if value is None else str(value).strip()


def child(parent, tag, value=None, **attributes):
    element = ET.SubElement(parent, f"{{{NS}}}{tag}", attributes)
    if value is not None:
        element.text = value
        element.attrib.update(LANG)
    return element


@dataclass
class Record:
    identifier: str
    kind: str
    name: str
    properties: dict
    endpoints: list
    location: str


def read_records(sheet, kind):
    """Find the header like tableToOFN, without indexing missing columns as -1."""
    headers = None
    for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
        values = [text(value) for value in row]
        if headers is None:
            indices = [i for i, value in enumerate(values) if NAME.lower() in value.lower()]
            if not indices:
                continue
            headers = values
            name_index = indices[0]
            endpoint_indices = [i for i, value in enumerate(headers)
                                if ENDPOINT.lower() in value.lower()]
            required = {CLASSES: 0, TROPES: 1, RELATIONSHIPS: 2}[kind]
            if len(endpoint_indices) != required:
                raise ValueError(f"{sheet.title}: expected {required} '{ENDPOINT}' columns")
            continue
        if not values[name_index]:
            continue
        endpoints = [values[i] for i in endpoint_indices]
        # tableToOFN also skips properties/relationships without endpoints.
        if any(not value for value in endpoints):
            warnings.warn(f"{sheet.title}, row {number}: skipping row with missing endpoint")
            continue
        properties = {header: value for i, (header, value) in enumerate(zip(headers, values))
                      if header and value and i != name_index and i not in endpoint_indices}
        prefix = {CLASSES: "class", TROPES: "trope", RELATIONSHIPS: "relation"}[kind]
        yield Record(f"id-{prefix}-{number}",
                     kind, values[name_index], properties, endpoints,
                     f"{sheet.title}, row {number}")
    if headers is None:
        raise ValueError(f"{sheet.title}: no '{NAME}' header found")


def property_value(record, name):
    return next((value for key, value in record.properties.items()
                 if name.lower() in key.lower()), "")


def hierarchy_edges(identifiers, relationships):
    """Use parent concepts first, then association direction where acyclic.

    Specialization points child -> parent in the model, but layout flows
    parent -> child. Conflicting/cyclic associations remain visible without
    being allowed to reverse the parent-concept hierarchy.
    """
    successors = {identifier: [] for identifier in identifiers}
    edges = []
    ordered = sorted(relationships, key=lambda link: link.attrib.get(f"{{{XSI}}}type") != "Specialization")
    for link in ordered:
        source, target = link.attrib["source"], link.attrib["target"]
        if link.attrib.get(f"{{{XSI}}}type") == "Specialization":
            source, target = target, source
        if source not in successors or target not in successors or source == target:
            continue
        pending, visited = [target], set()
        while pending:
            node = pending.pop()
            if node not in visited:
                visited.add(node)
                pending.extend(successors[node])
        if source in visited or target in successors[source]:
            continue
        successors[source].append(target)
        edges.append((source, target))
    return edges


def hierarchical_rows(identifiers, relationships, root_id=None):
    """Order real hierarchy levels by branch, wrapping each level at ten."""
    if root_id is not None:
        levels = [[root_id], [identifier for identifier in identifiers if identifier != root_id]]
    else:
        edges = hierarchy_edges(identifiers, relationships)
        successors = {identifier: [] for identifier in identifiers}
        predecessors = {identifier: set() for identifier in identifiers}
        for source, target in edges:
            successors[source].append(target)
            predecessors[target].add(source)
        # Preorder keeps siblings and their descendants together even when
        # the workbook interleaves unrelated branches.
        roots = [identifier for identifier in identifiers if not predecessors[identifier]]
        branch_order, visited = [], set()
        pending = list(reversed(roots))
        while pending:
            identifier = pending.pop()
            if identifier in visited:
                continue
            visited.add(identifier)
            branch_order.append(identifier)
            pending.extend(reversed(successors[identifier]))
        depth = {identifier: 0 for identifier in identifiers}
        remaining = set(identifiers)
        while remaining:
            ready = [identifier for identifier in branch_order
                     if identifier in remaining and not predecessors[identifier] & remaining]
            for identifier in ready:
                remaining.remove(identifier)
                for target in successors[identifier]:
                    depth[target] = max(depth[target], depth[identifier] + 1)
        levels = [[identifier for identifier in branch_order if depth[identifier] == level]
                  for level in sorted(set(depth.values()))]
    return [level[index:index + 10] for level in levels for index in range(0, len(level), 10)]


def add_hierarchical_view(diagrams, view_id, name, identifiers, relationships, root_id=None,
                          layout_relationships=None):
    """Draw bounded rows and route long links through nearby clear corridors."""
    view = child(diagrams, "view", identifier=view_id, **{f"{{{XSI}}}type": "Diagram"})
    child(view, "name", name)
    layout_relationships = relationships if layout_relationships is None else layout_relationships
    rows = hierarchical_rows(identifiers, layout_relationships, root_id)
    columns = max((len(row) for row in rows), default=0)
    cells = {identifier: (row_index, (columns - len(row)) // 2 + column)
             for row_index, row in enumerate(rows) for column, identifier in enumerate(row)}

    def reserve(lanes, start, end):
        start, end = sorted((start, end))
        for index, intervals in enumerate(lanes):
            if all(end < left or start > right for left, right in intervals):
                intervals.append((start, end))
                return index
        lanes.append([(start, end)])
        return len(lanes) - 1

    # Horizontal lanes can be shared when their spans do not overlap.
    routes = {}
    uses = {(identifier, side): [] for identifier in identifiers for side in ("source", "target")}
    for link in relationships:
        identifier = link.attrib["identifier"]
        source, target = link.attrib["source"], link.attrib["target"]
        source_row, _ = cells[source]
        target_row, _ = cells[target]
        routes[identifier] = {"needs_corridor": source_row + 1 != target_row}
        uses[(source, "source")].append(link)
        uses[(target, "target")].append(link)
    width = max(180, max((len(value) + 1 for value in uses.values()), default=0))
    column_x = [40 + column * (width + 60) for column in range(columns)]
    node_x = {}
    children = {identifier: [] for identifier in identifiers}
    for source, target in hierarchy_edges(identifiers, layout_relationships):
        children[source].append(target)
    if root_id is not None:
        children[root_id] = [identifier for identifier in identifiers if identifier != root_id]
    # Work upward from the leaves: place each parent over its children instead
    # of independently centering every row. Keep rows ordered and boxes apart.
    for row in reversed(rows):
        desired = []
        for index, identifier in enumerate(row):
            descendants = [node_x[target] for target in children[identifier] if target in node_x]
            fallback = 40 + ((columns - len(row)) * (width + 60)) // 2 + index * (width + 60)
            desired.append((min(descendants) + max(descendants)) // 2 if descendants else fallback)
        placed = []
        for x in desired:
            placed.append(max(x, placed[-1] + width + 60) if placed else x)
        overflow = max(0, placed[-1] - column_x[-1])
        placed = [x - overflow for x in placed]
        for index in range(len(placed) - 2, -1, -1):
            placed[index] = min(placed[index], placed[index + 1] - width - 60)
        shift = max(0, 40 - placed[0])
        node_x.update({identifier: x + shift for identifier, x in zip(row, placed)})
    right_edge = max((x + width for x in node_x.values()), default=40)
    ports = {}
    for (identifier, side), links in uses.items():
        other = "target" if side == "source" else "source"
        links.sort(key=lambda link: (cells[link.attrib[other]][1], cells[link.attrib[other]][0]))
        for index, link in enumerate(links, 1):
            ports[(link.attrib["identifier"], side)] = (
                node_x[identifier] + index * width // (len(links) + 1))

    # A long or backward link must pass the boxes on intermediate rows.
    # Search all open column gaps and use the nearest one with a free vertical
    # lane. A path goes outside the diagram only if these gaps are occupied.
    corridor_use = {}
    row_boxes = [[(node_x[identifier], node_x[identifier] + width)
                  for identifier in row] for row in rows]
    candidates = list(range(16, right_edge + 12 * (len(relationships) + 2), 12))
    for link in sorted(relationships, key=lambda item: (
            abs(cells[item.attrib["target"]][0] - cells[item.attrib["source"]][0]),
            item.attrib["identifier"])):
        identifier = link.attrib["identifier"]
        route = routes[identifier]
        if not route["needs_corridor"]:
            continue
        source_band = cells[link.attrib["source"]][0] + 1
        target_band = cells[link.attrib["target"]][0]
        first, last = sorted((source_band, target_band))
        blocked = [box for row in row_boxes[first:last] for box in row]
        sx, tx = ports[(identifier, "source")], ports[(identifier, "target")]
        for gx in sorted(candidates, key=lambda x: (abs(x - sx) + abs(x - tx),
                                                     abs(x - tx), x)):
            if any(left - 12 <= gx <= right + 12 for left, right in blocked):
                continue
            if any(not (last <= start or first >= end)
                   for start, end in corridor_use.get(gx, [])):
                continue
            corridor_use.setdefault(gx, []).append((first, last))
            route["x"] = gx
            break
        else:
            raise ValueError(f"No clear route for association {identifier}")

    bands = [[] for _ in range(len(rows) + 1)]
    for link in relationships:
        identifier = link.attrib["identifier"]
        route = routes[identifier]
        source_band = cells[link.attrib["source"]][0] + 1
        target_band = cells[link.attrib["target"]][0]
        sx, tx = ports[(identifier, "source")], ports[(identifier, "target")]
        if not route["needs_corridor"]:
            route["source_lane"] = route["target_lane"] = reserve(bands[source_band], sx, tx)
        else:
            gx = route["x"]
            route["source_lane"] = reserve(bands[source_band], sx, gx)
            route["target_lane"] = reserve(bands[target_band], tx, gx)
    band_y, row_y, y = [], [], 40
    for index, lanes in enumerate(bands):
        band_y.append(y)
        y += max(60, 12 * (len(lanes) + 2))
        if index < len(rows):
            row_y.append(y)
            y += 80
    references = {identifier: f"{view_id}-node-{identifier}" for identifier in identifiers}
    for identifier in identifiers:
        row, column = cells[identifier]
        child(view, "node", identifier=references[identifier], elementRef=identifier,
              x=str(node_x[identifier]), y=str(row_y[row]), w=str(width), h="80",
              **{f"{{{XSI}}}type": "Element"})
    for link in relationships:
        identifier = link.attrib["identifier"]
        source, target = link.attrib["source"], link.attrib["target"]
        sr, tr = cells[source][0], cells[target][0]
        sx, tx = ports[(identifier, "source")], ports[(identifier, "target")]
        route = routes[identifier]
        sy = band_y[sr + 1] + 12 * (route["source_lane"] + 1)
        ty = band_y[tr] + 12 * (route["target_lane"] + 1)
        points = [(sx, row_y[sr] + 80), (sx, sy)]
        if route["needs_corridor"]:
            points.extend([(route["x"], sy), (route["x"], ty)])
        points.extend([(tx, ty), (tx, row_y[tr])])
        connection = child(view, "connection", identifier=f"{view_id}-connection-{identifier}",
                           relationshipRef=identifier, source=references[source], target=references[target],
                           **{f"{{{XSI}}}type": "Relationship"})
        for px, py in points:
            child(connection, "bendpoint", x=str(px), y=str(py))


def add_default_views(root, elements, relationships):
    """Create a class overview and an owner-rooted view for each property set."""
    definitions = {definition.attrib["identifier"]: definition.findtext(f"{{{NS}}}name", "").lower()
                   for definition in root.findall(f"{{{NS}}}propertyDefinitions/{{{NS}}}propertyDefinition")}

    def element_type(element):
        for prop in element.findall(f"{{{NS}}}properties/{{{NS}}}property"):
            if definitions.get(prop.attrib["propertyDefinitionRef"]) == TYPE.lower():
                return prop.findtext(f"{{{NS}}}value", "").strip().lower()
        return ""

    classes = {element.attrib["identifier"]: element for element in elements
               if element_type(element) in ("typ objektu", "typ subjektu")}
    tropes = {element.attrib["identifier"] for element in elements if element_type(element) == "typ vlastnosti"}
    diagrams = child(child(root, "views"), "diagrams")
    associations = [link for link in relationships
                    if link.attrib[f"{{{XSI}}}type"] == "Association"
                    and link.attrib["source"] in classes and link.attrib["target"] in classes]
    layout_links = [link for link in relationships
                    if link.attrib[f"{{{XSI}}}type"] in ("Association", "Specialization")
                    and link.attrib["source"] in classes and link.attrib["target"] in classes]
    add_hierarchical_view(diagrams, "id-class-overview", "Subjekty a objekty práva",
                          list(classes), associations, layout_relationships=layout_links)
    for owner, element in classes.items():
        properties = set()
        for link in relationships:
            if link.attrib[f"{{{XSI}}}type"] not in ("Composition", "Association"):
                continue
            source, target = link.attrib["source"], link.attrib["target"]
            if source == owner and target in tropes:
                properties.add(target)
            elif target == owner and source in tropes:
                properties.add(source)
        if not properties:
            continue
        identifiers = [owner] + [element.attrib["identifier"] for element in elements
                                 if element.attrib["identifier"] in properties]
        selected = set(identifiers)
        links = [link for link in relationships
                 if link.attrib["source"] in selected and link.attrib["target"] in selected]
        name = element.findtext(f"{{{NS}}}name", owner)
        add_hierarchical_view(diagrams, f"id-properties-{owner}", f"{name} – vlastnosti",
                              identifiers, links, root_id=owner)


def convert(input_path, output_path, with_view=False):
    """Write exchange XML only after all workbook references have been resolved."""
    if Path(input_path).resolve() == Path(output_path).resolve():
        raise ValueError("Input and output paths must differ")
    workbook = openpyxl.load_workbook(input_path, data_only=True, read_only=True)
    try:
        for sheet in (VOCABULARY, CLASSES, TROPES, RELATIONSHIPS):
            if sheet not in workbook:
                raise ValueError(f"Missing worksheet: {sheet}")
        vocabulary_rows = list(workbook[VOCABULARY].iter_rows(values_only=True))
        name = text(vocabulary_rows[0][1]) if vocabulary_rows and len(vocabulary_rows[0]) > 1 else ""
        if not name:
            raise ValueError(f"{VOCABULARY}!B1 must contain the vocabulary name")
        vocabulary_properties = {}
        for row in vocabulary_rows[1:]:
            if len(row) > 1 and text(row[0]) and text(row[1]):
                vocabulary_properties[text(row[0])] = text(row[1])
        # The table reader uses B2 for the description regardless of its label.
        if len(vocabulary_rows) > 1 and len(vocabulary_rows[1]) > 1:
            description = text(vocabulary_rows[1][1])
            if description:
                vocabulary_properties["Popis"] = description
        records = [record for kind in (CLASSES, TROPES, RELATIONSHIPS)
                   for record in read_records(workbook[kind], kind)]
    finally:
        workbook.close()

    ET.register_namespace("", NS)
    ET.register_namespace("xsi", XSI)
    root = ET.Element(f"{{{NS}}}model", {
        "identifier": "id-vocabulary",
        f"{{{XSI}}}schemaLocation": f"{NS} https://www.opengroup.org/xsd/archimate/3.1/archimate3_Model.xsd",
    })
    child(root, "name", name)
    definitions = {}

    def add_properties(element, properties):
        if not properties:
            return
        container = child(element, "properties")
        for key, value in properties.items():
            if key not in definitions:
                definitions[key] = f"id-property-{len(definitions) + 1}"
            prop = child(container, "property", propertyDefinitionRef=definitions[key])
            child(prop, "value", value)

    add_properties(root, vocabulary_properties)
    if any(CATALOGUE.lower() in key.lower() for key in vocabulary_properties):
        warnings.warn("Catalogue address is preserved in XML, but archiToOFN.py ignores it")
    elements = child(root, "elements")
    relationships = child(root, "relationships")
    by_name = {}
    by_iri = {}
    nodes = {}

    def index(record):
        by_name.setdefault(record.name.lower(), []).append(record)
        iri = property_value(record, IRI)
        if iri:
            by_iri.setdefault(iri, []).append(record)

    for record in records:
        index(record)

    def resolve(reference, location):
        reference = reference.strip()
        matches = by_iri.get(reference) or by_name.get(reference.lower(), [])
        if not matches:
            reference = unquote(reference)
            matches = by_iri.get(reference) or by_name.get(reference.lower(), [])
        if len(matches) == 1:
            return matches[0]
        if matches:
            raise ValueError(f"{location}: ambiguous reference {reference!r}; use a unique IRI")
        if not re.match(r"^[A-Za-z][A-Za-z0-9+.-]*:", reference):
            raise ValueError(f"{location}: unknown concept {reference!r}")
        record = Record(f"id-external-{len(nodes) + 1}", CLASSES, reference,
                        {IRI: reference}, [], location)
        index(record)
        node = child(elements, "element", identifier=record.identifier,
                     **{f"{{{XSI}}}type": "BusinessObject"})
        child(node, "name", reference)
        add_properties(node, record.properties)
        nodes[record.identifier] = node
        warnings.warn(f"{location}: external IRI {reference!r} requires a placeholder; "
                      "archiToOFN.py may also export it as a term")
        return record

    for record in records:
        properties = record.properties.copy()
        if record.kind != RELATIONSHIPS:
            node = child(elements, "element", identifier=record.identifier,
                         **{f"{{{XSI}}}type": "BusinessObject"})
            if record.kind == TROPES:
                ofn_type = "Typ vlastnosti"
            else:
                ofn_type = {"subjekt práva": "Typ subjektu", "objekt práva": "Typ objektu"}.get(
                    property_value(record, TYPE).lower(), "")
                if not ofn_type:
                    warnings.warn(f"{record.location}: no subject/object type; archiToOFN.py "
                                  "will import this class as a generic term")
            # Keep Typ first: the importer uses substring matching, which could
            # otherwise mistake 'Datový typ' or 'Typ obsahu údaje' for it.
            properties = {TYPE: ofn_type, **{key: value for key, value in properties.items()
                                          if key != TYPE}} if ofn_type else properties
        else:
            source, target = [resolve(value, record.location) for value in record.endpoints]
            node = child(relationships, "relationship", identifier=record.identifier,
                         source=source.identifier, target=target.identifier, isDirected="true",
                         **{f"{{{XSI}}}type": "Association"})
        child(node, "name", record.name)
        add_properties(node, properties)
        nodes[record.identifier] = node

    def link(kind, source, target):
        child(relationships, "relationship", identifier=f"id-link-{len(relationships) + 1}",
              source=source, target=target, **{f"{{{XSI}}}type": kind})

    # Associations must precede specializations: the OFN importer processes
    # relationships in document order and needs the association terms first.
    for record in records:
        if record.kind == TROPES:
            owner = resolve(record.endpoints[0], record.location)
            link("Composition", owner.identifier, record.identifier)
        for reference in property_value(record, PARENT).split(";"):
            if reference.strip():
                parent = resolve(reference, record.location)
                link("Specialization", record.identifier, parent.identifier)

    if not len(elements):
        root.remove(elements)
    if not len(relationships):
        root.remove(relationships)
    if definitions:
        container = child(root, "propertyDefinitions")
        for name, identifier in definitions.items():
            definition = child(container, "propertyDefinition", identifier=identifier, type="string")
            child(definition, "name", name)
    if with_view:
        root.set(f"{{{XSI}}}schemaLocation",
                 f"{NS} https://www.opengroup.org/xsd/archimate/3.1/archimate3_Diagram.xsd")
        add_default_views(root, elements, relationships)
    ET.indent(root, space="  ")
    ET.ElementTree(root).write(output_path, encoding="utf-8", xml_declaration=True)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("input", type=Path, help="Excel workbook (.xlsx)")
    parser.add_argument("output", type=Path, help="ArchiMate exchange XML (.xml)")
    parser.add_argument("--with-view", action="store_true",
                        help="include a class overview and per-class property views (top-down, at most 10 nodes per row)")
    args = parser.parse_args()
    try:
        convert(args.input, args.output, with_view=args.with_view)
    except (ValueError, OSError, KeyError, BadZipFile, InvalidFileException) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
